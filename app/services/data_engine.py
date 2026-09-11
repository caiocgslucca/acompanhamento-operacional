"""Motor analítico: Excel/CSV -> Polars -> Parquet -> DuckDB."""
from pathlib import Path
import hashlib, json
import duckdb
import polars as pl

class DataEngine:
    def __init__(self, root='data'):
        self.root=Path(root); self.root.mkdir(exist_ok=True)
        self.manifest=self.root/'manifest.json'

    def fingerprint(self, path: Path):
        s=path.stat(); return hashlib.sha256(f'{path.name}:{s.st_size}:{s.st_mtime_ns}'.encode()).hexdigest()

    def ingest(self, source: str, dataset: str):
        path=Path(source); marks=json.loads(self.manifest.read_text('utf-8')) if self.manifest.exists() else {}
        mark=self.fingerprint(path)
        if marks.get(dataset)==mark: return {'status':'unchanged','dataset':dataset}
        if path.suffix.lower()=='.csv': frame=pl.scan_csv(path, infer_schema_length=10000).collect()
        elif path.suffix.lower() in ('.xlsx','.xlsm'):
            from openpyxl import load_workbook
            book=load_workbook(path,read_only=True,data_only=True); sheet=book.active
            rows=sheet.iter_rows(values_only=True); columns=[str(v).strip() if v is not None else f'col_{i}' for i,v in enumerate(next(rows))]
            frame=pl.DataFrame(list(rows),schema=columns,orient='row',infer_schema_length=10000)
        else: raise ValueError('Formato não suportado. Use CSV ou XLSX.')
        out=self.root/f'{dataset}.parquet'; frame.write_parquet(out,compression='zstd',statistics=True)
        marks[dataset]=mark; self.manifest.write_text(json.dumps(marks,indent=2),encoding='utf-8')
        return {'status':'published','dataset':dataset,'rows':frame.height,'file':str(out)}

    def query(self, dataset: str, sql: str, params=None):
        path=str((self.root/f'{dataset}.parquet').resolve()).replace("'","''")
        con=duckdb.connect(':memory:'); con.execute(f"CREATE VIEW dados AS SELECT * FROM read_parquet('{path}')")
        return con.execute(sql,params or []).pl()
