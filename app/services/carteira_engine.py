import csv
import hashlib
import os
import re
import shutil
import sys
import time
import unicodedata
import zipfile
from pathlib import Path
from threading import Lock, get_ident
import polars as pl
from openpyxl import load_workbook
from app.services import config_store as store

_cache={}; _lock=Lock(); _source_locks={}; EXTENSIONS={'.xlsx','.xlsm','.csv'}
STAGING=store.DATA_DIR/'staging'/'sources'; STAGING.mkdir(parents=True,exist_ok=True)

class SourceAccessError(RuntimeError): pass

def invalidate_cache(source_name=None):
    with _lock:
        if source_name:_cache.pop(clean_name(source_name),None)
        else:_cache.clear()

def windows_copy(source,destination):
    """Cópia nativa simples, sem abrir o documento ou alterar o OneDrive."""
    if sys.platform!='win32': return False,None
    try:
        import ctypes
        ok=ctypes.windll.kernel32.CopyFileW(str(source),str(destination),False)
        return bool(ok),None if ok else ctypes.get_last_error()
    except Exception as exc:return False,str(exc)

def safe_snapshot(path,attempts=3):
    """Copia a origem sincronizada para disco local antes da leitura analítica."""
    key=hashlib.sha256(str(path.resolve()).lower().encode('utf-8')).hexdigest()[:18]
    source_stat=path.stat()
    # O destino também é versionado. Assim uma leitura em andamento nunca
    # bloqueia a publicação feita por outra atualização no Windows.
    version=f'{source_stat.st_size}.{source_stat.st_mtime_ns}.{os.getpid()}.{get_ident()}'
    target=STAGING/f'{key}.{version}{path.suffix.lower()}'
    # Atualizações manuais, automáticas e telas podem ler a mesma fonte ao
    # mesmo tempo. Cada execução precisa do próprio arquivo temporário.
    part=STAGING/f'{key}.{os.getpid()}.{get_ident()}.part'
    last_error=None
    for attempt in range(1,attempts+1):
        try:
            part.unlink(missing_ok=True)
            with path.open('rb') as source,part.open('wb') as destination:
                shutil.copyfileobj(source,destination,length=8*1024*1024)
                destination.flush(); os.fsync(destination.fileno())
            if part.stat().st_size!=path.stat().st_size: raise OSError('A cópia local ficou incompleta.')
            if path.suffix.lower() in ('.xlsx','.xlsm') and not zipfile.is_zipfile(part): raise ValueError('O arquivo Excel não é um XLSX válido ou ainda está sincronizando.')
            os.replace(part,target)
            store.logger.info('LOCAL_SNAPSHOT_READY | arquivo=%s | bytes=%s | tentativa=%s',path.name,target.stat().st_size,attempt)
            return target,False
        except (PermissionError,OSError,ValueError) as exc:
            last_error=exc
            store.logger.warning('SOURCE_READ_RETRY | arquivo=%s | tentativa=%s/%s | erro=%s',path.name,attempt,attempts,exc)
            if attempt==1:
                copied,copy_error=windows_copy(path,part)
                if copied:
                    try:
                        if part.stat().st_size!=path.stat().st_size: raise OSError('A cópia nativa ficou incompleta.')
                        if path.suffix.lower() in ('.xlsx','.xlsm') and not zipfile.is_zipfile(part): raise ValueError('O Excel ainda não está disponível localmente.')
                        os.replace(part,target); store.logger.info('WINDOWS_NATIVE_COPY_READY | arquivo=%s | bytes=%s',path.name,target.stat().st_size); return target,False
                    except Exception as native_validation_error:last_error=native_validation_error
                store.logger.warning('WINDOWS_NATIVE_COPY_UNAVAILABLE | arquivo=%s | erro=%s',path.name,copy_error)
            time.sleep(min(.35*attempt,1.4))
    part.unlink(missing_ok=True)
    previous=sorted(STAGING.glob(f'{key}.*{path.suffix.lower()}'),key=lambda item:item.stat().st_mtime_ns,reverse=True)
    if previous:
        store.logger.warning('SOURCE_SNAPSHOT_FALLBACK | arquivo=%s | motivo=%s',path.name,last_error)
        return previous[0],True
    raise SourceAccessError(f'O arquivo “{path.name}” está bloqueado pelo Excel, OneDrive ou por falta de permissão. Feche o arquivo, confirme que ele está disponível neste dispositivo e tente novamente.') from last_error

def clean_name(value):
    text=unicodedata.normalize('NFKD',str(value or '')).encode('ascii','ignore').decode().upper().strip()
    text=re.sub(r'\s+(NUMBER|VARCHAR2|DATE)$','',text); return re.sub(r'[^A-Z0-9_]+','_',text).strip('_')

def files_for(path_text):
    path=Path(path_text)
    if not path.exists(): raise FileNotFoundError(f'Caminho da fonte não encontrado: {path}')
    files=[path] if path.is_file() else sorted(p for p in path.rglob('*') if p.is_file() and p.suffix.lower() in EXTENSIONS and not p.name.startswith('~$'))
    if not files: raise FileNotFoundError(f'Nenhum arquivo XLSX, XLSM ou CSV localizado em: {path}')
    return files

def read_file(path):
    if path.suffix.lower()=='.csv':
        sample=path.read_text(encoding='utf-8-sig',errors='replace')[:4096]
        try: separator=csv.Sniffer().sniff(sample,delimiters=';,\t|').delimiter
        except csv.Error: separator=';'
        return pl.read_csv(path,separator=separator,try_parse_dates=True,infer_schema_length=10000,ignore_errors=True,encoding='utf8-lossy')
    book=load_workbook(path,read_only=True,data_only=True); sheet=book.active; rows=sheet.iter_rows(values_only=True)
    try: header=next(rows)
    except StopIteration: return pl.DataFrame()
    columns=[]; seen={}
    for i,value in enumerate(header):
        name=clean_name(value) or f'COL_{i+1}'; seen[name]=seen.get(name,0)+1; columns.append(name if seen[name]==1 else f'{name}_{seen[name]}')
    data=list(rows); book.close()
    if not data:return pl.DataFrame(schema=columns)
    series=[]
    for index,name in enumerate(columns):
        values=(row[index] if index<len(row) else None for row in data)
        series.append(pl.Series(name,values,dtype=pl.String,strict=False))
    return pl.DataFrame(series)

def load_source(source_name,force=False):
    key=clean_name(source_name)
    with _lock:source_lock=_source_locks.setdefault(key,Lock())
    with source_lock:return _load_source(source_name,force)

def _load_source(source_name,force=False):
    source=next((s for s in store.list_sources() if clean_name(s['name'])==clean_name(source_name)),None)
    if not source: raise FileNotFoundError(f'Cadastre uma fonte com o nome “{source_name}” em Configurações.')
    files=files_for(source['path']); signature=tuple((str(p),p.stat().st_size,p.stat().st_mtime_ns) for p in files)
    key=clean_name(source_name)
    with _lock:
        if not force and key in _cache and _cache[key][0]==signature:return _cache[key][1].clone(),source,False
    frames=[]; all_columns=[]
    stale_files=[]
    for path in files:
        local_path,stale=safe_snapshot(path); stale_files.append(path.name) if stale else None
        frame=read_file(local_path)
        if frame.height:
            frame=frame.rename({c:clean_name(c) for c in frame.columns}).with_columns(pl.lit(path.name).alias('_ARQUIVO_ORIGEM'))
            frames.append(frame); all_columns.extend(frame.columns)
    if not frames: raise ValueError(f'Os arquivos da fonte {source_name} não possuem dados.')
    columns=list(dict.fromkeys(all_columns)); aligned=[]
    for frame in frames:
        missing=[pl.lit(None).alias(c) for c in columns if c not in frame.columns]
        aligned.append(frame.with_columns(missing).select(columns))
    result=pl.concat(aligned,how='vertical_relaxed')
    with _lock: _cache[key]=(signature,result.clone())
    store.logger.info('SOURCE_CONSOLIDATED | fonte=%s | arquivos=%s | linhas=%s | colunas=%s | forçada=%s | snapshots_anteriores=%s',source_name,len(files),result.height,len(result.columns),force,len(stale_files))
    return result,source,bool(stale_files)

def col(frame,*names):
    for name in names:
        clean=clean_name(name)
        if clean in frame.columns:return clean
    return None

def text_expr(name): return pl.col(name).cast(pl.Utf8,strict=False).fill_null('').str.strip_chars()

AGING_SCHEMA=['MAIOR QUE D7','D7','D6','D5','D4','D3','D2','D1','D0']
CLASS_SCHEMA=['ZCHP','ZFER','ZFOR','ZLEO','ZMAD','ZMAQ','ZQUI','ZTAB','ZTRI']
SUPPLEMENT_SOURCE_NAMES=('A_EXTRAIR','AGUARDANDO_LIBERACAO','CARTEIRA_AGUARDANDO')

def natural_key(value):
    match=re.match(r'^\s*(\d+)',str(value))
    return (int(match.group(1)) if match else 999999,str(value))

def matrix_rows(frame,row_col,column_col,value_col,fixed_columns=None,fixed_rows=None,row_header='DS_ONDA'):
    grouped=(frame.filter((text_expr(row_col)!='')&(text_expr(column_col)!='')).with_columns(pl.col(value_col).cast(pl.Float64,strict=False).fill_null(0).alias('_VALOR')).group_by([row_col,column_col]).agg(pl.col('_VALOR').sum().alias('valor')))
    rows={}; columns=[]
    for item in grouped.iter_rows(named=True):
        row=str(item[row_col]); column=str(item[column_col]); columns.append(column); rows.setdefault(row,{})[column]=round(item['valor'] or 0)
    discovered=sorted(set(columns),key=natural_key)
    columns=list(dict.fromkeys([*(fixed_columns or []),*discovered]))
    output=[]
    row_labels=list(dict.fromkeys([*(fixed_rows or []),*sorted(rows,key=natural_key)]))
    for row in row_labels:
        row_values=rows.get(row,{})
        values={c:row_values.get(c,0) for c in columns}; output.append({'label':row,'values':values,'total':sum(values.values())})
    totals={c:sum(r['values'][c] for r in output) for c in columns}
    return {'row_header':row_header,'columns':columns,'rows':output,'totals':totals,'grand_total':sum(totals.values())}

def supplemental_sources():
    result=[]
    for source in store.list_sources():
        name=clean_name(source['name']);path_name=clean_name(Path(source['path']).name)
        if name in SUPPLEMENT_SOURCE_NAMES or 'A_EXTRAIR' in path_name:result.append(source)
    return result

def carteira_related_names():return {'CARTEIRA','ONDA',*(clean_name(s['name']) for s in supplemental_sources())}

def load_carteira_combined():
    carteira,source,stale=load_source('Carteira');frames=[carteira];used=[]
    carteira_files={str(path.resolve()).lower() for path in files_for(source['path'])}
    for extra in supplemental_sources():
        extra_files={str(path.resolve()).lower() for path in files_for(extra['path'])}
        if extra_files and extra_files.issubset(carteira_files):continue
        frame,_,extra_stale=load_source(extra['name']);frames.append(frame);used.append(extra['name']);stale=stale or extra_stale
    if len(frames)==1:return carteira,source,stale,used
    columns=list(dict.fromkeys(column for frame in frames for column in frame.columns));aligned=[]
    for frame in frames:
        aligned.append(frame.with_columns([pl.lit(None).alias(c) for c in columns if c not in frame.columns]).select(columns))
    combined=pl.concat(aligned,how='vertical_relaxed')
    store.logger.info('CARTEIRA_COMPLEMENTED | fontes=%s | linhas=%s',','.join(used),combined.height)
    return combined,source,stale,used

def build(filters=None):
    filters=filters or {}; carteira,_,_,_=load_carteira_combined(); onda,_,_=load_source('Onda')
    cd_onda=col(carteira,'CD_ONDA'); qtd=col(carteira,'QTD_PENDENTE'); aging=col(carteira,'DS'); classe=col(carteira,'CD_CLASSE'); rota=col(carteira,'CD_ROTA'); status=col(carteira,'STATUS_SEPARACAO'); empresa=col(carteira,'CD_EMPRESA','EMPRESA','FILIAL')
    onda_key=col(onda,'CD_ONDA'); ds_onda=col(onda,'DS_ONDA','DS_ONDA_VARCHAR2')
    missing=[name for name,value in [('CD_ONDA',cd_onda),('QTD_PENDENTE',qtd),('DS',aging),('CD_CLASSE',classe)] if not value]
    if not onda_key or not ds_onda: missing.append('Onda: CD_ONDA/DS_ONDA')
    if missing: raise ValueError('Colunas obrigatórias ausentes: '+', '.join(missing))
    join_key='_CD_ONDA_JOIN'
    label_col='_DS_ONDA_LABEL';mapping=onda.with_columns(text_expr(onda_key).alias(join_key),text_expr(ds_onda).alias(label_col)).select([join_key,label_col]).unique(subset=[join_key],keep='first')
    data=carteira.with_columns(text_expr(cd_onda).alias(join_key)).join(mapping,on=join_key,how='left')
    source_ds=col(carteira,'DS_ONDA','DS_ONDA_VARCHAR2')
    if source_ds:data=data.with_columns(pl.when(text_expr(source_ds)!='').then(text_expr(source_ds)).otherwise(pl.col(label_col).fill_null('')).alias(label_col))
    data=data.with_columns(pl.when(pl.col(label_col).fill_null('')=='').then(pl.lit('Onda não cadastrada')).otherwise(pl.col(label_col)).alias(label_col))
    # Linhas vazias, fórmulas apagadas e sobras de formatação do Excel não são
    # registros operacionais. Assim o indicador acompanha de fato o arquivo.
    data=data.filter((text_expr(cd_onda)!='')&(text_expr(qtd)!=''))
    fixed_wave_rows=sorted((v for v in mapping.select(pl.col(label_col).alias('v')).unique().get_column('v').to_list() if v),key=natural_key)
    filter_columns={'empresa':empresa,'classe':classe,'rota':rota,'status':status}
    aging_data=data
    options={}
    for key,column in filter_columns.items():
        if column:
            values=[v for v in data.select(text_expr(column).alias('v')).unique().get_column('v').to_list() if v]
            if key=='status':values=[v for v in values if clean_name(v) in ('PENDENTE','EM_SEPARACAO','AGUARDANDO_LIBERACAO')]
            options[key]=sorted(values)
            selected=filters.get(key) or []
            if isinstance(selected,str):selected=[selected]
            if selected:
                data=data.filter(text_expr(column).is_in(selected))
                if key!='classe':aging_data=aging_data.filter(text_expr(column).is_in(selected))
        else:options[key]=[]
    status_selected=bool(filters.get('status'))
    if status_selected:
        positive=pl.col(qtd).cast(pl.Float64,strict=False).fill_null(0)>0
        data=data.filter(positive);aging_data=aging_data.filter(positive)
    aging_columns=list(dict.fromkeys([*AGING_SCHEMA,*(v for v in data.select(text_expr(aging).alias('v')).unique().get_column('v').to_list() if v)]))
    class_columns=list(dict.fromkeys([*CLASS_SCHEMA,*(v for v in data.select(text_expr(classe).alias('v')).unique().get_column('v').to_list() if v)]))
    # A matriz por faixa representa exclusivamente a classe ZCHP, conforme a
    # regra operacional, mesmo quando outras classes forem escolhidas no filtro.
    aging_data=aging_data.filter(text_expr(classe)=='ZCHP')
    source=next(s for s in store.list_sources() if clean_name(s['name'])=='CARTEIRA')
    fixed_rows=None if status_selected else fixed_wave_rows
    return {'matrix_aging':matrix_rows(aging_data,label_col,aging,qtd,aging_columns,fixed_rows),'matrix_class':matrix_rows(data,label_col,classe,qtd,class_columns,fixed_rows),'filters':options,'meta':{'rows':data.height,'updated_at':source.get('last_run'),'status':source.get('status') or 'IDLE','status_message':source.get('status_message') or '','total_pending':round(data.select(pl.col(qtd).cast(pl.Float64,strict=False).fill_null(0).sum()).item() or 0)}}
